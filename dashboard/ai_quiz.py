import json

from django.conf import settings
from google import genai

from .models import DEFAULT_QUESTIONS_FALLBACK


def extract_text_from_file(django_file) -> str:
    name = django_file.name.lower()
    django_file.seek(0)

    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(django_file)
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    elif name.endswith(".docx"):
        import docx
        doc = docx.Document(django_file)
        return "\n".join(p.text for p in doc.paragraphs)

    elif name.endswith(".pptx"):
        from pptx import Presentation
        prs = Presentation(django_file)
        text_runs = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text_runs.append(shape.text)
        return "\n".join(text_runs)

    return ""


def generate_questions_with_ai(lesson_text: str, level_group: str, num_questions: int = 15):
    if not settings.GEMINI_API_KEY or not lesson_text.strip():
        print("AI generation skipped — missing API key or empty extracted text.")
        return DEFAULT_QUESTIONS_FALLBACK

    audience = "preschool children (ages 3-5)" if level_group == "Preschool" else "elementary school children (ages 6-12)"

    prompt = f"""You are writing a quiz for {audience} based on the lesson content below.

Create exactly {num_questions} multiple-choice questions. Each question must have
exactly 4 short answer options, with exactly one correct answer. Keep language
simple and age-appropriate for {audience}. Base every question directly on the
lesson content provided — do not invent facts not present in it.

Respond with ONLY valid JSON, no markdown formatting, no explanation, matching
this exact schema:

{{
  "questions": [
    {{"text": "question text", "options": ["opt1", "opt2", "opt3", "opt4"], "correct": 0}}
  ]
}}

"correct" is the zero-based index of the right answer in "options".

Lesson content:
\"\"\"
{lesson_text[:8000]}
\"\"\"
"""

    try:
        client = genai.Client(api_key=settings.GEMINI_API_KEY)
        response = client.models.generate_content(
            model="gemini-3.6-flash",
            contents=prompt,
        )
        raw = response.text.strip()
        raw = raw.removeprefix("```json").removeprefix("```").removesuffix("```").strip()

        print("RAW GEMINI RESPONSE:", raw[:500])  # temporary — remove once working

        parsed = json.loads(raw)
        questions = parsed.get("questions", [])

        valid = []
        for q in questions:
            if (
                isinstance(q.get("options"), list)
                and len(q["options"]) == 4
                and isinstance(q.get("correct"), int)
                and 0 <= q["correct"] < 4
                and q.get("text")
            ):
                valid.append(q)

        return valid if valid else DEFAULT_QUESTIONS_FALLBACK

    except Exception as e:
        import traceback
        print("AI question generation failed:", e)
        traceback.print_exc()  # temporary — remove once working
        return DEFAULT_QUESTIONS_FALLBACK